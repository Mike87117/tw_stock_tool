from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from tw_stock_tool.backtesting import walk_forward
from tw_stock_tool.cli import twstock_cli
from tw_stock_tool.reports.walk_forward_report import (
    build_walk_forward_report_data,
)


def _detail(errors: list[str]) -> pd.DataFrame:
    return pd.DataFrame([
        {"Window": i, "Strategy": "ma_cross", "Error": error,
         "Test Sharpe Ratio": 99. if error else 1.,
         "Test Total Return %": 99. if error else 5.,
         "Test CAGR %": 99. if error else 10.,
         "Test Max Drawdown %": -99. if error else -2.}
        for i, error in enumerate(errors, 1)
    ]).reindex(columns=walk_forward.WALK_FORWARD_COLUMNS)


def _summary(detail: pd.DataFrame) -> pd.Series:
    return walk_forward.build_summary(detail, "2330", "1y", "ma_cross", 4, 2, 2).iloc[0]


class WalkForwardFailureStatusTest(unittest.TestCase):
    def test_all_failed_evaluations_have_no_best_window_or_average(self):
        detail = _detail(["backtest failed", "backtest failed"])
        data = build_walk_forward_report_data(detail)
        self.assertEqual(data["Summary"]["Status"], "ERROR")
        self.assertEqual(data["Summary"]["Valid Rows"], 0)
        self.assertEqual(data["Summary"]["Error Rows"], 2)
        self.assertIsNone(data["Best Window"])
        self.assertNotIn("Average Test Total Return %", data["Summary"])
        self.assertEqual(len(data["Results"]), 2)
        summary = _summary(detail)
        self.assertEqual(summary["Status"], "ERROR")
        self.assertTrue(pd.isna(summary["Avg Test Total Return %"]))
        self.assertTrue(pd.isna(summary["Positive Test Windows %"]))

    def test_partial_failure_is_excluded_from_ranking_and_averages(self):
        detail = _detail(["backtest failed", ""])
        original = detail.copy(deep=True)
        data = build_walk_forward_report_data(detail)
        self.assertEqual(data["Summary"]["Status"], "PARTIAL")
        self.assertEqual(data["Best Window"]["Window"], 2)
        self.assertEqual(data["Summary"]["Average Test Total Return %"], 5.)
        self.assertEqual(data["Summary"]["Worst Test Max Drawdown %"], -2.)
        summary = _summary(detail)
        self.assertEqual(summary["Status"], "PARTIAL")
        self.assertEqual(summary["Avg Test Total Return %"], 5.)
        self.assertEqual(summary["Error Windows"], 1)
        self.assertIn("backtest failed", summary["Error"])
        pd.testing.assert_frame_equal(detail, original)

    def test_empty_results_are_errors(self):
        data = build_walk_forward_report_data(_detail([]))
        self.assertEqual(data["Summary"]["Status"], "ERROR")
        self.assertIsNone(data["Best Window"])
        self.assertIn("No walk-forward windows", data["Summary"]["Error"])
        self.assertEqual(_summary(_detail([]))["Status"], "ERROR")

    def test_nan_metrics_cannot_select_a_best_window(self):
        detail = _detail([""])
        detail["Test Sharpe Ratio"] = float("nan")
        data = build_walk_forward_report_data(detail)
        self.assertIsNone(data["Best Window"])

    def test_successful_run_and_legacy_frame_remain_supported(self):
        for detail in (_detail(["", ""]), _detail(["", ""]).drop(columns="Error")):
            with self.subTest(columns=detail.columns):
                data = build_walk_forward_report_data(detail)
                self.assertEqual(data["Summary"]["Status"], "OK")
                self.assertEqual(data["Summary"]["Error Rows"], 0)
                self.assertEqual(data["Best Window"]["Window"], 1)

    def test_real_invalid_parameter_grid_fails_unified_cli_and_exports_diagnostics(self):
        prices = pd.DataFrame({"Open": [100.] * 12, "High": [102.] * 12, "Low": [99.] * 12,
                               "Close": [100.] * 12, "Volume": [1000] * 12},
                              index=pd.date_range("2026-01-01", periods=12))
        with tempfile.TemporaryDirectory() as temporary:
            markdown = Path(temporary) / "report.md"
            excel = Path(temporary) / "report.xlsx"
            with patch.object(walk_forward, "analyze_stock", return_value=SimpleNamespace(signal_df=prices)), \
                 redirect_stdout(StringIO()) as output:
                status = twstock_cli.main([
                    "walk-forward", "--stock", "2330", "--strategy", "ma_cross",
                    "--train-days", "4", "--test-days", "2",
                    "--ma-short-windows", "5", "--ma-long-windows", "3",
                    "--output-md", str(markdown), "--output-excel", str(excel),
                ])
            self.assertEqual(status, 1)
            self.assertIn("No valid walk-forward windows", output.getvalue())
            self.assertIn("No valid ma_cross parameter combinations", output.getvalue())
            self.assertNotIn("Walk forward finished", output.getvalue())
            text = markdown.read_text(encoding="utf-8")
            self.assertIn("Status: ERROR", text)
            self.assertIn("No top walk-forward window found.", text)
            self.assertIn("No valid ma_cross parameter combinations", text)
            workbook = load_workbook(excel, read_only=True)
            try:
                self.assertEqual(list(workbook["Top Walk-Forward Window"].values), [])
                rows = dict(workbook["Summary"].values)
                self.assertEqual(rows["Status"], "ERROR")
                self.assertEqual(rows["Error Rows"], 4)
                self.assertEqual(len(list(workbook["Results"].values)), 5)
            finally:
                workbook.close()

    def test_both_cli_entrypoints_report_all_failure_and_empty_results(self):
        from tw_stock_tool.cli import walk_forward_report
        report_args = walk_forward_report._parse_args(["--stock", "2330", "--strategy", "ma_cross"])
        with patch("sys.argv", ["walk_forward", "--stock", "2330"]):
            legacy_args = walk_forward._parse_args()
        for module, args in ((walk_forward_report, report_args), (walk_forward, legacy_args)):
            for detail in (_detail(["backtest failed"]), _detail([])):
                with self.subTest(module=module.__name__, rows=len(detail)):
                    with patch.object(module, "_parse_args", return_value=args), \
                         patch.object(module, "run_walk_forward", return_value=detail), \
                         redirect_stdout(StringIO()) as output:
                        self.assertEqual(module.main(), 1)
                    self.assertIn("Error:", output.getvalue())
                    self.assertNotIn("Top Walk-Forward Strategy:", output.getvalue())

    def test_both_cli_entrypoints_preserve_partial_and_successful_results(self):
        from tw_stock_tool.cli import walk_forward_report
        report_args = walk_forward_report._parse_args(["--stock", "2330", "--strategy", "ma_cross"])
        with patch("sys.argv", ["walk_forward", "--stock", "2330"]):
            legacy_args = walk_forward._parse_args()
        for module, args in ((walk_forward_report, report_args), (walk_forward, legacy_args)):
            for errors in (["", "backtest failed"], ["", ""]):
                with self.subTest(module=module.__name__, errors=errors):
                    with patch.object(module, "_parse_args", return_value=args), \
                         patch.object(module, "run_walk_forward", return_value=_detail(errors)), \
                         redirect_stdout(StringIO()) as output:
                        self.assertIsNone(module.main())
                    self.assertEqual("Warning:" in output.getvalue(), bool(errors[1]))

    def test_legacy_excel_summary_preserves_failure_status_and_blank_metrics(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "legacy.xlsx"
            walk_forward.export_walk_forward_excel(_detail(["backtest failed"]), "2330", "1y", "ma_cross", 4, 2, 2, str(path))
            workbook = load_workbook(path, read_only=True)
            try:
                rows = list(workbook["Summary"].values)
                summary = dict(zip(rows[0], rows[1]))
                self.assertEqual(summary["Status"], "ERROR")
                self.assertIsNone(summary["Avg Test Total Return %"])
                self.assertEqual(summary["Error Windows"], 1)
                self.assertEqual(len(list(workbook["Errors"].values)), 2)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
